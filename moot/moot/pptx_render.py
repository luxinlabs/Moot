"""Render a styled deck to .pptx with python-pptx: themed backgrounds, decorative motifs, and drawn visuals
(flow diagrams, numbered steps, timelines, stat cards), plus speaker notes."""
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONTS = ["Inter", "Helvetica", "Georgia", "Montserrat", "Source Sans Pro", "Playfair Display"]
BACKGROUNDS = ["solid", "gradient", "glow", "grid", "dots", "bands"]
MOTIFS = ["none", "circles", "corner", "lines"]
VISUALS = ["none", "flow", "steps", "timeline", "stat_cards"]
DEFAULT_THEME = {"name": "Moot default", "background": "#0F1218", "text": "#F2F4F8", "accent": "#8B7FF5", "accent_2": "#3DDC97",
                 "muted": "#9AA3B2", "heading_font": "Inter", "body_font": "Inter", "background_style": "glow", "motif": "circles"}
W, H = Inches(13.333), Inches(7.5)
HEX = re.compile(r"#?[0-9a-fA-F]{6}")


def _hex(h):
    return str(h or "").strip().lstrip("#").upper()


def rgb(h):
    h = _hex(h)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def mix(a, b, t):
    """Blend hex colour a toward b by t (0..1)."""
    a, b = _hex(a), _hex(b)
    return "#" + "".join(f"{round(int(a[i:i + 2], 16) * (1 - t) + int(b[i:i + 2], 16) * t):02X}" for i in (0, 2, 4))


def luminance(h):
    h = _hex(h)
    if not HEX.fullmatch(h):
        return 0.0
    def ch(c):
        c = int(c, 16) / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(h[i:i + 2]) for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def clean_theme(t):
    t = {**DEFAULT_THEME, **(t or {})}
    for k in ("background", "text", "accent", "accent_2", "muted"):
        v = str(t.get(k) or "").strip()
        t[k] = "#" + _hex(v if HEX.fullmatch(v) else DEFAULT_THEME[k])
    for k in ("heading_font", "body_font"):
        if t[k] not in FONTS:
            t[k] = "Inter"
    if t["background_style"] not in BACKGROUNDS:
        t["background_style"] = "solid"
    if t["motif"] not in MOTIFS:
        t["motif"] = "none"
    return t


# ---------------------------------------------------------------- shape helpers
def _alpha(fill, opacity):
    clr = fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))


def shape(sl, kind, x, y, w, h, color, opacity=1.0, line=None, rot=0):
    s = sl.shapes.add_shape(kind, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if color:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        if opacity < 1:
            _alpha(s.fill, opacity)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line[0])
        s.line.width = Pt(line[1])
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if rot:
        s.rotation = rot
    return s


def text(sl, x, y, w, h, value, font, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = sl.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.04))
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value or ""
    r.font.name, r.font.size, r.font.bold, r.font.italic = font, Pt(size), bold, italic
    r.font.color.rgb = rgb(color)
    return box


def bullets(sl, x, y, w, h, items, t, size):
    box = sl.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items or []):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.45)
        d = p.add_run()
        d.text = "■  "
        d.font.size, d.font.name = Pt(size * 0.6), t["body_font"]
        d.font.color.rgb = rgb(t["accent"])
        r = p.add_run()
        r.text = item
        r.font.name, r.font.size = t["body_font"], Pt(size)
        r.font.color.rgb = rgb(t["text"])


# ---------------------------------------------------------------- backgrounds + motifs
def background(sl, t):
    bg, acc, acc2, fg = t["background"], t["accent"], t["accent_2"], t["text"]
    style = t["background_style"]
    fill = sl.background.fill
    if style == "gradient":
        fill.gradient()
        fill.gradient_angle = 35
        st = fill.gradient_stops
        st[0].color.rgb, st[0].position = rgb(bg), 0.0
        st[1].color.rgb, st[1].position = rgb(mix(bg, acc, 0.28)), 1.0
    else:
        fill.solid()
        fill.fore_color.rgb = rgb(bg)
    if style == "glow":
        shape(sl, MSO_SHAPE.OVAL, W * 0.62, -H * 0.45, W * 0.6, W * 0.6, acc, 0.09)
        shape(sl, MSO_SHAPE.OVAL, -W * 0.22, H * 0.55, W * 0.45, W * 0.45, acc2, 0.07)
    elif style == "grid":
        step = Inches(0.6)
        for i in range(1, int(W / step) + 1):
            shape(sl, MSO_SHAPE.RECTANGLE, i * step, 0, Inches(0.01), H, fg, 0.06)
        for j in range(1, int(H / step) + 1):
            shape(sl, MSO_SHAPE.RECTANGLE, 0, j * step, W, Inches(0.01), fg, 0.06)
    elif style == "dots":
        for i in range(24):
            for j in range(13):
                if (i + j) % 2 == 0:
                    shape(sl, MSO_SHAPE.OVAL, Inches(0.3 + i * 0.55), Inches(0.3 + j * 0.55), Inches(0.05), Inches(0.05), fg, 0.16)
    elif style == "bands":
        shape(sl, MSO_SHAPE.RECTANGLE, W * 0.55, -H * 0.3, Inches(2.2), H * 1.6, acc, 0.10, rot=20)
        shape(sl, MSO_SHAPE.RECTANGLE, W * 0.78, -H * 0.3, Inches(1.1), H * 1.6, acc2, 0.10, rot=20)


def motif(sl, t):
    acc, acc2 = t["accent"], t["accent_2"]
    m = t["motif"]
    if m == "circles":
        for k, (d, o) in enumerate(((3.6, 0.06), (2.4, 0.09), (1.2, 0.16))):
            shape(sl, MSO_SHAPE.OVAL, W - Inches(d * 0.75), H - Inches(d * 0.7), Inches(d), Inches(d), acc if k != 1 else acc2, o)
    elif m == "corner":
        shape(sl, MSO_SHAPE.RIGHT_TRIANGLE, W - Inches(1.6), 0, Inches(1.6), Inches(1.6), acc, 0.85, rot=180)
        shape(sl, MSO_SHAPE.RIGHT_TRIANGLE, W - Inches(0.9), 0, Inches(0.9), Inches(0.9), acc2, 0.9, rot=180)
    elif m == "lines":
        for k in range(3):
            shape(sl, MSO_SHAPE.RECTANGLE, Inches(0.7), H - Inches(1.05 - k * 0.14), Inches(2.4 - k * 0.6), Inches(0.05), acc if k != 1 else acc2, 0.9)


# ---------------------------------------------------------------- visuals
def visual(sl, kind, items, t, x, y, w, h):
    items = [i for i in (items or []) if i and i.strip()][:6]
    if kind == "none" or not items:
        return
    acc, acc2, fg, mut = t["accent"], t["accent_2"], t["text"], t["muted"]
    font, n = t["body_font"], len(items)
    if kind == "flow":
        gap = Inches(0.45)
        bw = (w - gap * (n - 1)) / n
        for i, it in enumerate(items):
            bx = x + i * (bw + gap)
            b = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, bw, h, acc, 0.16, line=(acc, 1.5))
            b.adjustments[0] = 0.18
            text(sl, bx + Inches(0.08), y, bw - Inches(0.16), h, it, font, 15 if n > 4 else 17, fg, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                shape(sl, MSO_SHAPE.RIGHT_ARROW, bx + bw + Inches(0.08), y + h / 2 - Inches(0.14), gap - Inches(0.16), Inches(0.28), acc2)
    elif kind == "steps":
        cw = w / n
        for i, it in enumerate(items):
            cx = x + i * cw + cw / 2
            d = Inches(0.72)
            c = shape(sl, MSO_SHAPE.OVAL, cx - d / 2, y, d, d, acc)
            tf = c.text_frame
            tf.text = str(i + 1)
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(22), True, t["heading_font"]
            r.font.color.rgb = rgb(t["background"])
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            if i < n - 1:
                shape(sl, MSO_SHAPE.RECTANGLE, cx + d / 2 + Inches(0.1), y + d / 2 - Inches(0.02), cw - d - Inches(0.2), Inches(0.04), mut, 0.6)
            text(sl, cx - cw / 2 + Inches(0.1), y + d + Inches(0.12), cw - Inches(0.2), h - d - Inches(0.12), it, font, 15, fg, align=PP_ALIGN.CENTER)
    elif kind == "timeline":
        ly = y + h * 0.42
        shape(sl, MSO_SHAPE.RECTANGLE, x, ly, w, Inches(0.05), acc, 0.7)
        cw = w / n
        for i, it in enumerate(items):
            cx = x + i * cw + cw / 2
            shape(sl, MSO_SHAPE.OVAL, cx - Inches(0.15), ly - Inches(0.13), Inches(0.3), Inches(0.3), acc2 if i == n - 1 else acc)
            ty = y if i % 2 == 0 else ly + Inches(0.3)
            text(sl, cx - cw / 2 + Inches(0.05), ty, cw - Inches(0.1), h * 0.4, it, font, 14, fg, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.BOTTOM if i % 2 == 0 else MSO_ANCHOR.TOP)
    elif kind == "stat_cards":
        gap = Inches(0.35)
        cw = (w - gap * (n - 1)) / n
        for i, it in enumerate(items):
            val, _, lab = it.partition("|")
            cx = x + i * (cw + gap)
            c = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, y, cw, h, fg, 0.07, line=(acc, 1))
            c.adjustments[0] = 0.12
            text(sl, cx + Inches(0.2), y + Inches(0.15), cw - Inches(0.4), h * 0.55, val.strip(), t["heading_font"], 34, acc if i % 2 == 0 else acc2, bold=True)
            text(sl, cx + Inches(0.2), y + h * 0.58, cw - Inches(0.4), h * 0.4, lab.strip(), font, 14, mut)


# ---------------------------------------------------------------- slides
def render(deck, path, product=""):
    t = clean_theme(deck.get("theme"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    slides = deck.get("slides", [])
    hf, bf, fg, acc, mut = t["heading_font"], t["body_font"], t["text"], t["accent"], t["muted"]
    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        background(sl, t)
        motif(sl, t)
        layout, vis = s.get("layout", "bullets"), s.get("visual", "none")
        has_vis = vis in VISUALS and vis != "none" and bool(s.get("visual_items"))
        title = s.get("title", "")
        vx, vy, vw, vh = Inches(1.0), Inches(4.45), Inches(11.3), Inches(1.75)
        if layout in ("title", "statement", "closing"):
            top = Inches(1.35) if has_vis else Inches(2.2)
            shape(sl, MSO_SHAPE.RECTANGLE, Inches(0.9), top, Inches(1.3), Inches(0.12), acc)
            text(sl, Inches(0.9), top + Inches(0.3), Inches(11.2), Inches(1.9), title, hf, 50 if layout == "title" else 40, fg, bold=True)
            if s.get("subtitle"):
                text(sl, Inches(0.9), top + Inches(2.15), Inches(11.2), Inches(0.9), s["subtitle"], bf, 22, mut)
            if layout == "closing" and s.get("bullets") and not has_vis:
                bullets(sl, Inches(0.9), Inches(5.4), Inches(11.2), Inches(1.4), s["bullets"][:3], t, 19)
        else:
            shape(sl, MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.72), Inches(0.12), Inches(0.95), acc)
            text(sl, Inches(1.0), Inches(0.55), Inches(11.4), Inches(1.3), title, hf, 32, fg, bold=True)
            if layout == "big_number":
                text(sl, Inches(1.0), Inches(1.95), Inches(11.4), Inches(1.9), s.get("stat_value", ""), hf, 88, acc, bold=True)
                text(sl, Inches(1.0), Inches(3.75), Inches(11.4), Inches(0.7), s.get("stat_label", ""), bf, 24, mut)
                if not has_vis:
                    bullets(sl, Inches(1.0), Inches(4.7), Inches(11.4), Inches(2.0), s.get("bullets", [])[:3], t, 19)
                vy = Inches(4.75)
            elif layout == "two_column":
                bullets(sl, Inches(1.0), Inches(2.1), Inches(5.5), Inches(4.6), s.get("bullets", []), t, 20)
                shape(sl, MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(2.1), Inches(0.03), Inches(4.2), mut, 0.4)
                bullets(sl, Inches(7.1), Inches(2.1), Inches(5.5), Inches(4.6), s.get("bullets_right", []), t, 20)
                has_vis = False
            elif layout == "quote":
                text(sl, Inches(0.9), Inches(1.4), Inches(2), Inches(2), "“", hf, 160, acc, bold=True)
                text(sl, Inches(1.9), Inches(2.4), Inches(10.2), Inches(2.8), s.get("subtitle") or " ".join(s.get("bullets", [])[:1]), hf, 32, fg, italic=True)
                has_vis = False
            else:
                y0 = Inches(1.8)
                if s.get("subtitle"):
                    text(sl, Inches(1.0), y0, Inches(11.4), Inches(0.6), s["subtitle"], bf, 20, mut)
                    y0 += Inches(0.6)
                bl = s.get("bullets", [])[:3] if has_vis else s.get("bullets", [])
                bullets(sl, Inches(1.0), y0 + Inches(0.1), Inches(11.4), Inches(2.2 if has_vis else 4.6), bl, t, 20 if has_vis else 23)
        if has_vis:
            visual(sl, vis, s.get("visual_items"), t, vx, vy, vw, vh)
        text(sl, Inches(0.7), Inches(6.95), Inches(8), Inches(0.35), product or deck.get("title", ""), bf, 11, mut)
        text(sl, Inches(11.6), Inches(6.95), Inches(1.0), Inches(0.35), f"{i + 1}/{len(slides)}", bf, 11, mut, align=PP_ALIGN.RIGHT)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(path)
    return path
