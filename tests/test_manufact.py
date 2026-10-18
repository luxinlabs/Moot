import asyncio
import base64
import copy
import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path

os.environ["MCP_USE_ANONYMIZED_TELEMETRY"] = "false"
from mcp_use import MCPClient
from pptx import Presentation

from fixture_studio import FixtureStudio
from moot.contracts import MootError
from moot.external_skills import VENDOR, load_external_skills
from moot.manufact_server import WorkspaceAuth, build_server
from moot.studio import load_skills, skills_block, skills_version
from moot.studio_agent import scripted_demo, unpack
from moot.studio_service import StudioService


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.studio = FixtureStudio(self.tmp.name)
        self.service = StudioService(self.studio, self.tmp.name)
        self.job = self.service.create("Deck", "https://moot.example/", candidates=["fixture"])

    def tearDown(self):
        self.studio.release.set()
        self.service.close()
        self.studio.store.db.close()
        self.tmp.cleanup()

    def wait(self):
        for _ in range(100):
            job = self.service.get(self.job["job_id"])
            if job["status"] in ("COMPLETED", "FAILED"):
                return job
            time.sleep(.02)
        self.fail("Fixture did not finish")

    def test_duplicate_answers_start_one_trial(self):
        self.studio.release.clear()
        jid = self.job["job_id"]
        self.service.answer_and_start(jid, {"budget": 3})
        self.service.answer_and_start(jid, {"budget": 3.0})
        with self.assertRaises(MootError):
            self.service.answer_and_start(jid, {"budget": 4})
        self.studio.release.set()
        self.assertEqual(self.wait()["status"], "COMPLETED")
        self.assertEqual(self.studio.run_count, 1)

    def test_invalid_answers_do_not_start_work(self):
        for answers in [{"budget": -1}, {"budget": "NaN"}, {"budget": float("inf")},
                        {"budget": "oops"}, {"audience": "unknown"}, {"unknown": "x"}]:
            with self.subTest(answers=answers), self.assertRaises(MootError):
                self.service.answer_and_start(self.job["job_id"], answers)
        self.assertEqual(self.studio.run_count, 0)
        self.assertEqual(self.service.get(self.job["job_id"])["status"], "QUESTIONS")

    def test_failure_becomes_terminal(self):
        self.studio.fail = True
        self.service.answer_and_start(self.job["job_id"], {})
        self.assertEqual(self.wait()["status"], "FAILED")

    def test_artifact_integrity_and_path_boundary(self):
        jid = self.job["job_id"]
        with self.assertRaises(MootError):
            self.service.artifact(jid)
        self.service.answer_and_start(jid, {})
        self.wait()
        manifest, data = self.service.artifact(jid)
        self.assertEqual(manifest["sha256"], hashlib.sha256(data).hexdigest())
        self.assertEqual(len(Presentation(io.BytesIO(data)).slides), 2)
        job = self.studio.get(jid)
        job["final"]["pptx"] = "../outside.pptx"
        self.studio.store.save_job(job)
        with self.assertRaises(MootError):
            self.service.artifact(jid)

    def test_web_uses_shared_service(self):
        from moot import web
        prior = (web.STUDIO, web.DECKS)
        web.STUDIO, web.DECKS = self.studio, self.service
        try:
            # WSGI client exercises Flask dispatch without importing its optional CLI test runner.
            from werkzeug.test import Client
            from werkzeug.wrappers import Response
            client = Client(web.app, Response)
            jid = self.job["job_id"]
            response = client.post(f"/api/studio/jobs/{jid}/answers", json={"budget": 3})
            self.assertEqual(response.status_code, 200)
            self.wait()
            response = client.get(f"/api/studio/jobs/{jid}")
            self.assertEqual(response.json["status"], "COMPLETED")
            self.assertEqual(self.service.review(jid)["winner"], "fixture")
        finally:
            web.STUDIO, web.DECKS = prior


class SkillTests(unittest.TestCase):
    def test_external_skills_reach_prompt_and_fingerprint(self):
        skills = load_skills()
        self.assertEqual(len([s for s in skills if "provenance" in s]), 2)
        block = skills_block(skills)
        self.assertIn("Pitch Deck Advisor", block)
        self.assertIn("deckgen slide design", block)
        self.assertIn("Moot adapter rules", block)
        altered = copy.deepcopy(skills)
        altered[-1]["provenance"]["commit"] = "0" * 40
        self.assertNotEqual(skills_version(skills), skills_version(altered))

    def test_hash_tampering_fails_closed(self):
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "vendor"
            shutil.copytree(VENDOR, dest)
            (dest / "deckgen" / "SKILL.md").write_text("tampered")
            with self.assertRaises(ValueError):
                load_external_skills(dest)


class ProtocolTests(unittest.IsolatedAsyncioTestCase):
    async def test_authenticated_http_transport(self):
        import httpx
        with tempfile.TemporaryDirectory() as folder:
            studio = FixtureStudio(folder)
            service = StudioService(studio, folder)
            server = build_server(service, token="test-secret")
            headers = {"Accept": "application/json, text/event-stream"}
            init = {"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {
                "protocolVersion": "2025-06-18", "capabilities": {},
                "clientInfo": {"name": "Moot HTTP integration test", "version": "1"}}}
            try:
                async with server.session_manager.run():
                    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="http://127.0.0.1:8000") as client:
                        denied = await client.post("/mcp", headers=headers, json=init)
                        self.assertEqual(denied.status_code, 401)
                        wrong = await client.post("/mcp", headers={**headers, "Authorization": "Bearer wrong"}, json=init)
                        self.assertEqual(wrong.status_code, 401)
                        headers["Authorization"] = "Bearer test-secret"
                        result = await client.post("/mcp", headers=headers, json=init)
                        self.assertEqual(result.status_code, 200, result.text)
                        self.assertIn("Moot Deck Studio", result.text)
                        if result.headers.get("mcp-session-id"):
                            headers["mcp-session-id"] = result.headers["mcp-session-id"]
                        headers["mcp-protocol-version"] = "2025-06-18"
                        await client.post("/mcp", headers=headers, json={"jsonrpc": "2.0", "method": "notifications/initialized"})
                        listed = await client.post("/mcp", headers=headers, json={"jsonrpc": "2.0", "id": 2, "method": "tools/list"})
                        self.assertEqual(listed.status_code, 200)
                        self.assertIn("moot_studio_create", listed.text)
                        self.assertIn("readOnlyHint", listed.text)
            finally:
                service.close()
                studio.store.db.close()

    async def test_auth(self):
        auth = WorkspaceAuth("test-secret")
        self.assertIsNone(await auth.verify_token("wrong"))
        self.assertEqual((await auth.verify_token("test-secret")).claims["sub"], "moot-workspace")

    async def test_real_stdio_round_trip(self):
        script = Path(__file__).with_name("fixture_server.py")
        root = Path(__file__).resolve().parents[1]
        client = MCPClient.from_dict({"mcpServers": {"moot": {
            "command": sys.executable, "args": [str(script)],
            "env": {"PYTHONPATH": str(root), "MCP_USE_ANONYMIZED_TELEMETRY": "false"}}}})
        try:
            with tempfile.TemporaryDirectory() as folder:
                output = Path(folder) / "demo.pptx"
                evidence = await scripted_demo(client, "TEST FIXTURE deck", "https://moot.example/", output, timeout=30)
                self.assertTrue(output.is_file())
                self.assertTrue(output.with_suffix(".evidence.json").is_file())
                session = client.get_session("moot")
                missing = await session.call_tool(name="moot_studio_get", arguments={"job_id": "missing"})
                self.assertTrue(missing.isError)
                # Exercise the binary MCP resource as well as the tool-only fallback.
                resource = await session.read_resource(evidence["artifact"]["resource_uri"])
                binary = base64.b64decode(resource.contents[0].blob)
                self.assertEqual(binary, output.read_bytes())
                self.assertEqual(len(Presentation(io.BytesIO(binary)).slides), 2)
        finally:
            await client.close_all_sessions()


if __name__ == "__main__":
    unittest.main()
